from __future__ import annotations

import re
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document


ROOT = Path(r"C:\Users\admin\Downloads\FCAJ-workshop-trungthanh")
SRC_DOCX = ROOT / "Nguyễn Trung Thành - Phiếu tiến độ.docx"
OUT_DIR = Path(r"C:\Users\admin\source\repos\Talent-Graph-AI\_outputs")
OUT_DOCX = OUT_DIR / "trungthanh-phieu-tien-do-sua-ngay-worklog.docx"


def fmt_date(value: datetime) -> str:
    return value.strftime("%d-%m-%Y")


def parse_dates_from_markdown(path: Path) -> list[datetime]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    values: list[datetime] = []
    for raw in re.findall(r"\b\d{1,2}/\d{1,2}/\d{4}\b", text):
        values.append(datetime.strptime(raw, "%d/%m/%Y"))
    return values


def collect_ranges() -> dict[int, tuple[str, str]]:
    content = ROOT / "content" / "1-Worklog"
    ranges: dict[int, tuple[str, str]] = {}

    for week_dir in sorted(content.glob("1.*-Week*")):
        match = re.match(r"1\.(\d+)-Week", week_dir.name)
        if not match:
            continue
        week = int(match.group(1))

        vi_path = week_dir / "_index.vi.md"
        en_path = week_dir / "_index.md"
        dates = parse_dates_from_markdown(vi_path if vi_path.exists() else en_path)

        # The Vietnamese Week 12 page is truncated in the provided folder; the
        # English page and Worklog overview carry the complete final-week span.
        if week == 12 and en_path.exists():
            en_dates = parse_dates_from_markdown(en_path)
            if len(en_dates) > len(dates):
                dates = en_dates

        if dates:
            ranges[week] = (fmt_date(min(dates)), fmt_date(max(dates)))

    return ranges


def replace_cell_text_preserve_style(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    style = paragraph.style
    alignment = paragraph.alignment

    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)

    for run in list(paragraph.runs):
        paragraph._element.remove(run._element)

    paragraph.style = style
    paragraph.alignment = alignment
    lines = text.split("\n")
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        paragraph.add_run(line)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SRC_DOCX, OUT_DOCX)

    ranges = collect_ranges()
    missing = sorted(set(range(1, 13)) - set(ranges))
    if missing:
        raise RuntimeError(f"Missing worklog dates for weeks: {missing}")

    document = Document(OUT_DOCX)
    changed: list[tuple[int, str]] = []

    for table in document.tables:
        for row in table.rows[1:]:
            week_text = row.cells[0].text
            match = re.search(r"\b(\d{1,2})\b", week_text)
            if not match:
                continue
            week = int(match.group(1))
            if week not in ranges:
                continue
            start, end = ranges[week]
            replace_cell_text_preserve_style(row.cells[1], f"{start}/\n {end}")
            changed.append((week, f"{start} - {end}"))

    if sorted(week for week, _ in changed) != list(range(1, 13)):
        raise RuntimeError(f"Updated unexpected week set: {changed}")

    document.save(OUT_DOCX)
    print(str(OUT_DOCX))
    for week, span in sorted(changed):
        print(f"Week {week}: {span}")


if __name__ == "__main__":
    main()
